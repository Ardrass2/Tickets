import os
import sys
from os import path
import random
import csv
import openpyxl
import sqlite3

from PyQt5.QtWidgets import QMainWindow, QApplication, QWidget, QMessageBox, \
    QTableWidgetItem, QHeaderView, QInputDialog, QButtonGroup, QFileDialog
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QIcon

from window_design.main_window import Ui_MainWindow as Main
from window_design.progress import Ui_Form as Progress
from window_design.terms import Ui_MainWindow as Terms
from window_design.tests import Ui_Form as Tests
from window_design.table import Ui_Form as Tables
from window_design.add_terms import Ui_Form as AddTerm


class EmptyValueError(Exception):
    def __init__(self, text):
        # Ошибка пустого значения
        self.msg = QMessageBox()
        self.msg.setIcon(QMessageBox.Warning)
        self.msg.setText(text)
        self.msg.setWindowTitle("Ошибка")
        self.msg.exec_()


class IdenticalDataError(Exception):
    def __init__(self, text):
        # Ошибка повторяющегося элемента значения
        self.msg = QMessageBox()
        self.msg.setIcon(QMessageBox.Critical)
        self.msg.setText(text)
        self.msg.setWindowTitle("Ошибка")
        self.msg.exec_()


class FewTermensError(Exception):
    def __init__(self, text, more_text=""):
        # Ошибка, связанная с количеством терминов
        self.msg = QMessageBox()
        self.msg.setText(text)
        self.msg.setInformativeText(more_text)
        self.msg.setWindowTitle("Мало терминов")
        self.msg.exec_()


class NoCategoryError(Exception):
    # Ошибка отсутствия категории
    def __init__(self):
        self.msg = QMessageBox()
        self.msg.setText("Не найдено категорий!")
        self.msg.setInformativeText("Чтобы продолжить добавьте категорию.")
        self.msg.setWindowTitle("Нет категорий")
        self.msg.exec_()


class MainWindow(QMainWindow, Main):
    # Главное окно
    def __init__(self):
        super().__init__()
        self.category = None
        self.setWindowIcon(QIcon("img/icon.png"))
        self.setupUi(self)
        self.stats.clicked.connect(self.change_window)
        self.db_button.clicked.connect(self.change_window)
        self.terms_button.clicked.connect(self.change_window)
        self.test_button.clicked.connect(self.change_window)
        # если не найдена база данных, создаем аналагичную
        if not path.exists('termens.db'):
            db = sqlite3.connect('termens.db')
            cur = db.cursor()
            cur.execute("""CREATE TABLE termen (
                        id         INTEGER PRIMARY KEY AUTOINCREMENT
                                           NOT NULL
                                           UNIQUE,
                        Terms      STRING  UNIQUE
                                           NOT NULL,
                        definition TEXT    NOT NULL
                                           UNIQUE,
                        [right]    STRING  DEFAULT ('s'),
                        [group]    INTEGER REFERENCES category (id) ON DELETE CASCADE
                                                                    ON UPDATE CASCADE
                                           NOT NULL
                    );""")
            cur.execute("""CREATE TABLE category (
                        id    INTEGER PRIMARY KEY AUTOINCREMENT
                                      NOT NULL
                                      UNIQUE,
                        title STRING  UNIQUE
                                      NOT NULL
                                );""")
            db.close()
        # с этим курсором мы будем работать во всех классах
        self.db = sqlite3.connect('termens.db')
        self.cur = self.db.cursor()
        self.window = None

    def change_window(self):
        # Смена окна
        button = QApplication.instance().sender()
        if button is self.test_button:
            try:
                if len(self.cur.execute("SELECT id FROM termen").fetchall()) >= 4:
                    categories = list(map(lambda x: str(x[0]), self.cur.execute("SELECT title FROM category")))
                    self.category, ok_pressed = QInputDialog.getItem(self, "Термины из категории...",
                                                                     "Выберите категорию", categories)
                    if ok_pressed:
                        if len(self.cur.execute(f"""SELECT terms FROM termen
                                            WHERE [group] = (SELECT id FROM category
                                            WHERE title = '{self.category}')""").fetchall()) >= 4:
                            self.window = TestWindow()
                            self.window.show()
                            self.hide()
                        else:
                            raise FewTermensError("У вас меньше 4 терминов",
                                                  more_text="Нужно как минимум 4 термина в 1 категории")
                else:
                    raise FewTermensError("У вас меньше 4 терминов",
                                          more_text="Такое количество лучше учить во вкладке Термины")
            except FewTermensError:
                pass
            except NoCategoryError:
                pass
        if button is self.db_button:
            self.window = Table()
            self.window.show()
            self.hide()
        if button is self.terms_button:
            try:
                if len(self.cur.execute("SELECT id FROM termen").fetchall()) >= 2:
                    categories = list(map(lambda x: str(x[0]), self.cur.execute("SELECT title FROM category")))
                    self.category, ok_pressed = QInputDialog.getItem(self, "Термины из категории...",
                                                                     "Выберите категорию", categories)
                    if ok_pressed:
                        if (len(self.cur.execute(f"""SELECT terms FROM termen
                                                    WHERE [group] = (SELECT id FROM category
                                                    WHERE title = '{self.category}')""").fetchall()) >= 2):
                            self.window = TermsWindow()
                            self.window.show()
                            self.hide()
                        else:
                            FewTermensError("В этой категории меньше 2 терминов!")
                else:
                    raise FewTermensError("У вас меньше 2 терминов")
            except FewTermensError:
                pass
        if button is self.stats:
            self.window = ProgressWindow()
            self.window.show()
            self.hide()


class TermsWindow(QMainWindow, Terms):
    # Окно с карточками
    def __init__(self):
        super().__init__()
        self.setupUi(self)
        self.label_2.setText('')
        self.setWindowIcon(QIcon("img/icon.png"))

        self.id = None
        self.right_term = None
        self.rights = None

        self.random_text()

        self.pushButton.clicked.connect(back_to_main)
        self.give_answ.clicked.connect(self.check_answer)
        self.next_bt.clicked.connect(self.next_text)

    def keyPressEvent(self, event):
        # обработка нажатий на клавиатуру
        if event.key() == Qt.Key_Escape:
            back_to_main()
        if event.key() == Qt.Key_Enter:
            self.check_answer()

    def random_text(self):
        # Выбор рандомного определения
        base = ex.cur.execute(f"""SELECT * FROM termen
                                WHERE [group] = (SELECT id FROM category
                                WHERE title = '{ex.category}')""").fetchall()
        num_of_text = random.randint(0, len(base) - 1)
        # Сохраним все данные, чтобы больше не образащаться за ними в БД
        self.textBrowser.setText(str(base[num_of_text][2]))
        self.id = str(base[num_of_text][0])
        self.right_term = str(base[num_of_text][1])
        self.rights = str(base[num_of_text][3])

    def next_text(self):
        # Замена определения на другое
        try:
            if self.label_2.text():
                old = self.textBrowser.toPlainText()
                while self.textBrowser.toPlainText() == old:
                    self.random_text()
                self.label_2.setText("")
                self.label_2.setStyleSheet("")
                self.label_4.setText("")
                self.label_4.setStyleSheet("")
                self.answ_edit.setText("")
                self.answ_edit.setStyleSheet("")
                self.answ_edit.setReadOnly(False)
            else:
                raise EmptyValueError("Необходимо ответить на вопрос!")
        except EmptyValueError:
            pass

    def check_answer(self):
        # Проверка ответа
        try:
            if not self.answ_edit.isReadOnly():
                if all([self.answ_edit.text()]):
                    if len(self.rights) == 21:
                        self.rights = self.rights[:-1]
                    if self.answ_edit.text().lower() == str(self.right_term).lower():
                        self.answ_edit.setStyleSheet("background: green;"
                                                     "color: white;")
                        self.label_2.setText("ПРАВИЛЬНЫЙ ОТВЕТ!")
                        if len(self.rights) >= 2:
                            ex.cur.execute(f"""UPDATE termen
                                    SET right = '{self.rights[0] + "1" + self.rights[1:]}'
                                    WHERE id = {self.id}""")
                        else:
                            ex.cur.execute(f"""UPDATE termen
                                    SET right = 's1'
                                    WHERE id = {self.id}""")
                    else:
                        self.answ_edit.setStyleSheet("background: red;"
                                                     "color: black;")
                        self.label_2.setText("Правильный ответ:")
                        self.label_4.setText(self.right_term)
                        self.label_4.setStyleSheet("background: green;"
                                                   "color: white;")
                        if len(self.rights) >= 2:
                            ex.cur.execute(f"""UPDATE termen
                                            SET right = '{self.rights[0] + "0" + self.rights[1:]}'
                                            WHERE id = {self.id}""")
                        else:
                            ex.cur.execute(f"""UPDATE termen
                                            SET right = 's0'
                                            WHERE id = {self.id}""")
                    ex.db.commit()
                    self.answ_edit.setReadOnly(True)
                else:
                    raise EmptyValueError("Вы ничего не ввели")
        except EmptyValueError:
            pass


class Table(QWidget, Tables):
    # Окно с таблицей, отображающий информацию из БД
    def __init__(self):
        super().__init__()
        self.setupUi(self)
        self.setWindowIcon(QIcon("img/icon.png"))
        self.tableWidget.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.tableWidget.horizontalHeader().setSectionResizeMode(
            0, QHeaderView.Fixed)
        self.back.clicked.connect(back_to_main)

        self.categories = list(map(lambda x: str(x[0]), ex.cur.execute("SELECT title FROM category")))
        self.result = None
        self.window = None
        self.msg = None

        self.filling_table()
        self.add.clicked.connect(self.add_action)
        self.clear_btn.clicked.connect(self.clear_all)
        self.del_btn.clicked.connect(self.delete_term)
        self.load_btn.clicked.connect(self.select_table)
        self.pushButton.clicked.connect(self.add_category)
        self.del_catergory_button.clicked.connect(self.delete_category)
        self.update_combobox()
        self.comboBox.activated.connect(self.filling_table)

    def keyPressEvent(self, event):
        # Обработка нажатий
        if event.key() == Qt.Key_Escape:
            back_to_main()
        if event.key() == Qt.Key_Delete:
            self.delete_term()

    def delete_category(self):
        try:
            if not self.categories:
                raise EmptyValueError("Не найдено категорий!")
            category, ok_pressed = QInputDialog.getItem(self, "Удалить категорию",
                                                        "Выберите категорию", self.categories)
            if ok_pressed:
                msg = QMessageBox.question(self, "Все так плохо?",
                                           f"Вы действительно хотите удалить катергорию - '{category}'?\n"
                                           "При удалении категории все термины в нем удалятся!",
                                           QMessageBox.Yes, QMessageBox.No)
                if msg == QMessageBox.Yes:
                    ex.cur.execute(f"""DELETE FROM termen
                                    WHERE [group] = (SELECT id FROM category
                                                    WHERE title = '{category}')""")
                    ex.cur.execute(f"""DELETE FROM category
                                    WHERE title = '{category}'""")
                    ex.db.commit()
                    self.categories.remove(category)
                    self.update_combobox()
                    self.filling_table()
        except EmptyValueError:
            pass

    def delete_term(self):
        # Удаление термина
        try:
            if not self.categories:
                raise EmptyValueError("Не найдено категорий!")
            category, ok_pressed = QInputDialog.getItem(self, "Добавление терминов", "Выберите категорию",
                                                        self.categories, 1, False)
            if ok_pressed:
                termens = set(map(lambda x: str(x[0]), ex.cur.execute(f"""SELECT terms FROM termen
                                                                WHERE [group]= 
                                                                (SELECT id FROM category
                                                                WHERE title='{category}')""").fetchall()))
                if termens:
                    term, ok = QInputDialog.getItem(self, "Удаление термина",
                                                    "Выберите термин:", termens,
                                                    editable=False)
                    if ok:
                        msg = QMessageBox.question(self, 'Все так плохо?',
                                                   f"""Вы действительно хотите удалить термин - "{term}"?""",
                                                   QMessageBox.Yes, QMessageBox.No)
                        if msg == QMessageBox.Yes:
                            ex.cur.execute(f"""DELETE FROM termen
                                        WHERE terms = '{term}'""")
                            ex.db.commit()
                            self.filling_table()
                else:
                    raise FewTermensError("Не найдено ни одного термина в данной категории!")
        except FewTermensError:
            pass
        except EmptyValueError:
            pass

    def add_action(self):
        # Открытие окна обновления термина
        try:
            if not self.categories:
                raise EmptyValueError("Сначала добавьте категории")
            category, ok_pressed = QInputDialog.getItem(self, "Добавление терминов", "Выберите категорию",
                                                        self.categories, 1, False)
            if ok_pressed:
                self.window = AddTermens(category)
                self.window.show()
                self.filling_table()
        except EmptyValueError:
            pass

    def add_category(self):
        # Добавление категорий
        try:
            category, ok_pressed = QInputDialog.getText(self, "Категории",
                                                        "Введите название категории")
            if category not in [self.comboBox.itemText(i) for i in range(self.comboBox.count())]:
                if ok_pressed:
                    if category:
                        ex.cur.execute(f"INSERT INTO category(title) VALUES('{category}')")
                        self.categories.append(category)
                        ex.db.commit()
                        self.update_combobox()
                    else:
                        raise EmptyValueError("Вы ничего не ввели")
            else:
                raise IdenticalDataError("Такая категория уже есть")
        except EmptyValueError:
            pass
        except IdenticalDataError:
            pass

    def filling_table(self):
        # Заполение таблицы данными из БД
        if self.comboBox.currentText() != "Все термины":
            self.result = ex.cur.execute(f"""SELECT terms, definition FROM termen
                                        WHERE [group] = (
                                        SELECT id FROM category
                                        WHERE title = '{self.comboBox.currentText()}')""").fetchall()
        else:
            self.result = ex.cur.execute("SELECT terms, definition FROM termen").fetchall()
        if self.result:
            self.tableWidget.setRowCount(len(self.result))
            self.tableWidget.setColumnCount(len(self.result[0]))
            for i, elem in enumerate(self.result):
                for j, val in enumerate(elem):
                    self.tableWidget.setItem(i, j, QTableWidgetItem(str(val)))
        else:
            while self.tableWidget.rowCount() > 0:
                self.tableWidget.removeRow(0)

    def update_combobox(self):
        # Обновление данных в комбобоксе
        all_items = [self.comboBox.itemText(i) for i in range(self.comboBox.count())]
        for elem in list(self.categories):
            if elem not in all_items:
                self.comboBox.addItem(elem)
        if len(list(self.categories)) + 1 < len(all_items):
            for index, elem in enumerate(all_items):
                if elem not in self.categories and elem != "Все термины":
                    self.comboBox.removeItem(index)

    def xlsx_to_csv(self, file):
        # Перевод из xlsx формата в csv для удобного добавления
        worksheet = list(openpyxl.load_workbook(filename=file, data_only=True))
        table = open("output.csv", "w", encoding="utf-8", newline='')
        sheet = worksheet[0]
        writer = csv.writer(table, delimiter=';', quotechar='"')
        for row in sheet.iter_rows(values_only=True):
            if any(row):
                row = [str(x) for x in row if x is not None]
                writer.writerow(row)
        table.close()

    def add_table_to_table(self, file, category):
        # Добавление таблицы из вне
        try:
            if file:
                with open(file, encoding="utf8") as csvfile:
                    reader = csv.reader(csvfile, delimiter=';', quotechar='"')
                    for elem in reader:
                        ex.cur.execute(f"""INSERT INTO termen(terms, definition, [group]) VALUES('{elem[0]}',
                         '{elem[1]}', (SELECT id FROM category WHERE title = '{category}'))""")
                    ex.db.commit()
                    self.filling_table()
        except sqlite3.IntegrityError:
            msg = IdenticalDataError("Какой-то термин или определение уже есть в списке!\n"
                                     "Мы добавили только те, которые остутсвуют в базе.")
        finally:
            if path.exists("output.csv"):
                os.remove("output.csv")

    def select_table(self):
        # Выбор таблицы
        try:
            if not self.categories:
                raise FewTermensError("Сначала добавьте категорию!")
            category, ok_pressed = QInputDialog.getItem(self, "Добавление терминов", "Выберите категорию",
                                                        self.categories, 1, False)
            if ok_pressed:
                file = QFileDialog.getOpenFileName(self, "Выбирите таблицу с терминами", ".",
                                                   """Microsoft Excel Files (*.xlsx);;Primitive Table (*.csv)""")[0]
                if file.endswith(".xlsx"):
                    self.xlsx_to_csv(file)
                    self.add_table_to_table("output.csv", category)
                else:
                    self.add_table_to_table(file, category)
        except FewTermensError:
            pass

    def clear_all(self):
        # Удаление всех терминов
        try:
            if len(ex.cur.execute("SELECT id FROM category").fetchall()) != 0:
                msg = QMessageBox.question(self, 'Все так плохо?', "Вы действительно хотите удалить ВСЁ?",
                                           QMessageBox.Yes, QMessageBox.No)
                if msg == QMessageBox.Yes:
                    ex.cur.execute("DELETE FROM termen")
                    ex.cur.execute("DELETE FROM category")
                    self.categories = list()
                    ex.db.commit()
                    self.update_combobox()
                    self.filling_table()
            else:
                raise FewTermensError("Терминов и так нет, кого удалять-то?")
        except FewTermensError:
            pass


class AddTermens(QWidget, AddTerm):
    # Окно с добавлением терминов
    def __init__(self, category):
        super().__init__()
        self.setupUi(self)
        self.category = category
        self.setWindowIcon(QIcon("img/icon.png"))
        self.pushButton.clicked.connect(self.next)
        self.pushButton_2.clicked.connect(self.clear_term)
        self.term = ''
        self.definition = ''

    def keyPressEvent(self, event):
        # обработка нажатий
        if event == Qt.Key_Enter:
            self.next()
        if event == Qt.Key_Escape:
            self.clear_term()

    def next(self):
        # переключение на следующий этап
        try:
            if "Термин" in self.label.text():
                if self.plainTextEdit.toPlainText():
                    if self.plainTextEdit.toPlainText() in \
                            list(map(lambda x: str(x[0]), ex.cur.execute("SELECT terms FROM termen").fetchall())):
                        raise IdenticalDataError("Этот термин уже существует!")
                    else:
                        self.term = self.plainTextEdit.toPlainText()
                        self.label.setText("<html><head/><body><p align=\"center\">Определение</p></body></html>")
                        self.plainTextEdit.setPlainText('')
                else:
                    raise EmptyValueError("Вы ничего не ввели!")
            else:
                if self.plainTextEdit.toPlainText():
                    if self.plainTextEdit.toPlainText() in \
                            list(map(lambda x: str(x[0]), ex.cur.execute("SELECT definition FROM termen").fetchall())):
                        raise IdenticalDataError("К этому понятию уже существует термин!")
                    else:
                        self.definition = self.plainTextEdit.toPlainText()
                        self.add_to_db()
                        self.hide()
                else:
                    raise EmptyValueError("Вы ничего не ввели!")
        except EmptyValueError:
            pass
        except IdenticalDataError:
            pass

    def clear_term(self):
        # закрытие окна
        self.close()

    def add_to_db(self):
        # добавление в БД и заполнение таблицы
        ex.cur.execute(f"""INSERT INTO termen(Terms, definition, [group]) VALUES('{self.term}', '{self.definition}',
                                                                                (SELECT id FROM category
                                                                                WHERE title = '{self.category}'))""")
        ex.db.commit()
        ex.window.filling_table()


class TestWindow(QWidget, Tests):
    # Окно с тестиками)
    def __init__(self):
        super().__init__()
        self.setupUi(self)
        self.setWindowIcon(QIcon("img/icon.png"))
        self.label_4.setText("")
        self.group = QButtonGroup()
        self.group.addButton(self.f_ans, 1)
        self.group.addButton(self.s_ans, 2)
        self.group.addButton(self.trd_ans, 3)
        self.group.addButton(self.frth_ans, 4)

        self.id = None
        self.right_term = None
        self.rights = None
        self.random_text()
        self.random_termens()

        self.pushButton.clicked.connect(self.next)
        self.back.clicked.connect(back_to_main)
        self.group.buttonClicked.connect(self.check_answer)

    def keyPressEvent(self, event):
        # Обработка нажатий
        if event.key() == Qt.Key_Escape:
            back_to_main()
        if event.key() == Qt.Key_Enter:
            self.next()

    def random_text(self):
        # Рандомный выбор определения
        base = ex.cur.execute(f"""SELECT terms, definition, [right] FROM termen
                                WHERE [group] = (SELECT id FROM category
                                                WHERE title = '{ex.category}')""").fetchall()
        current_text = self.textBrowser.toPlainText()
        # Делаем так, чтобы термины были различны
        while current_text == self.textBrowser.toPlainText():
            num_of_text = random.randint(0, len(base) - 1)
            # По аналогии с TermsWindow сразу достаем все нам необходимое
            self.textBrowser.setText(str(base[num_of_text][1]))
            self.right_term = str(base[num_of_text][0])
            self.rights = str(base[num_of_text][2])

    def check_answer(self, button):
        # проверка ответа
        if self.rights == 21:
            self.rights = self.rights[:-1]
        if button.text() == self.right_term:
            self.label_4.setText("ВЕРНО")
            ex.cur.execute(f"""UPDATE termen
                            SET right = '{self.rights[0] + "1" + self.rights[1:]}'
                            WHERE terms = '{self.right_term}'""")
            button.setStyleSheet("""background-color: green;
                                color: white;""")

        else:
            self.label_4.setText("НЕВЕРНО")
            ex.cur.execute(f"""UPDATE termen
                            SET right = '{self.rights[0] + "0" + self.rights[1:]}'
                            WHERE terms = '{self.right_term}'""")
            button.setStyleSheet("""background-color: red;
                                color: white;""")
            for btn in self.group.buttons():
                if btn.text() == self.right_term:
                    btn.setStyleSheet("""background-color: green;
                                color: white;""")
        self.action_with_radiobutton(False)
        ex.db.commit()

    def random_termens(self):
        # Добавление 3 рандомных и 1 правильного термина на radiobutton
        termens = list(map(lambda x: str(x[0]), ex.cur.execute(f"""SELECT terms FROM termen
                                                               WHERE [group] = (SELECT id FROM category
                                                                                WHERE title = '{ex.category}')""")))

        termens.remove(self.right_term)
        temp_pos = random.sample(termens, 3)
        temp_pos.append(self.right_term)
        random_pos = random.sample(temp_pos, 4)
        for i, btn in enumerate(self.group.buttons()):
            btn.setText(random_pos[i])

    def action_with_radiobutton(self, flag):
        # установка доступности нажатия на radiobutton
        for btn in self.group.buttons():
            btn.setEnabled(flag)

    def next(self):
        # переключение на следующий термин
        try:
            if not self.group.button(1).isEnabled():
                self.random_text()
                self.random_termens()
                self.action_with_radiobutton(True)
                for btn in self.group.buttons():
                    btn.setStyleSheet("""background-color: #A7C4D4;
                                    color: #F6E8DA;""")
                self.label_4.setText("")
            else:
                raise EmptyValueError("Вы ничего не выбрали")
        except EmptyValueError:
            pass


class ProgressWindow(QWidget, Progress):
    # Окно статистики
    def __init__(self):
        super().__init__()
        self.setupUi(self)
        self.setWindowTitle("Статистика")
        self.progressBar.hide()
        self.setWindowIcon(QIcon("img/icon.png"))
        self.pushButton.clicked.connect(back_to_main)
        for elem in ex.cur.execute(f"""SELECT title FROM category"""):
            self.combo_category.addItem(str(elem[0]))
        self.combo_category.activated.connect(self.update_termens)
        self.combo_term.activated.connect(self.show_score)

    def update_termens(self):
        self.combo_term.clear()
        self.combo_term.addItem("Выберите термин")
        termens = ex.cur.execute(f"""SELECT terms FROM termen
                                    WHERE [group] = 
                                    (SELECT id FROM category
                                    WHERE title = '{self.combo_category.currentText()}')""").fetchall()
        for elem in termens:
            self.combo_term.addItem(str(elem[0]))
        self.progressBar.hide()

    def keyPressEvent(self, event):
        # Обработка нажатий
        if event.key() == Qt.Key_Escape:
            back_to_main()

    def show_score(self):
        # вывод на progressbar процента правильного ответа из последних 20 ответом
        if self.combo_term.currentText() != "Выберите термин":
            self.progressBar.show()
            list_of_answer = str(ex.cur.execute(f"""SELECT [right] FROM termen
                                        WHERE terms = '{self.combo_term.currentText()}'""").fetchall()[0][0])
            right = 0
            for letter in list_of_answer:
                if letter == '1':
                    right += 1
            if len(list_of_answer) - 1 != 0:
                percent = int((right / (len(list_of_answer) - 1)) * 100)
                self.progressBar.setValue(percent)
            else:
                self.progressBar.setValue(0)
        else:
            self.progressBar.hide()


def back_to_main():
    # функция возврата на главный экран
    ex.window.hide()
    ex.show()


def except_hook(cls, exception, traceback):
    # Отлов ошибок
    sys.__excepthook__(cls, exception, traceback)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    ex = MainWindow()
    ex.show()
    sys.excepthook = except_hook
    sys.exit(app.exec_())
